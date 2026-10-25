"""The Monday deck and the app's site data (UX v2): week names, the default week, the summary
slide, wording edits, exceptions, links, skip groups and the run log."""

from __future__ import annotations

import glob
import importlib.util
import json
from pathlib import Path

import pytest

from brightloop.collect.capture import Snapshot
from brightloop.collect.run import snapshot_ref
from brightloop.contract import check_contract
from brightloop.deck import (
    build_slides,
    is_wording_only,
    meeting_label,
    render_html,
    week_range,
    week_view,
    write_deck,
)
from brightloop.explain import skip_group
from brightloop.pipeline import run_pipeline

ROOT = Path(__file__).resolve().parents[1]
RECORDINGS = ROOT / "recordings"
FIXTURES = ROOT / "tests" / "fixtures"


def load_script(name: str):
    spec = importlib.util.spec_from_file_location(name, ROOT / "scripts" / f"{name}.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def recorded(week: str) -> tuple[dict, dict, dict]:
    folder = RECORDINGS / "weeks" / week
    return tuple(
        json.loads((folder / name).read_text(encoding="utf-8"))
        for name in ("cases.json", "result.json", "run.json")
    )


def view(week: str) -> dict:
    return week_view(*recorded(week))


def page(v: dict, title: str) -> dict:
    return next(s for s in v["slides"] if s["title"] == title)


# ---- Week names and the default week (spec 3.2, 3.4) -------------------------------------------


def test_weeks_are_named_after_their_monday_meeting():
    assert meeting_label("2026-W43") == "Monday, Oct 26"
    assert meeting_label("2026-W43", short=True) == "Mon, Oct 26"
    assert week_range("2026-W43") == "Oct 19–25"
    assert week_range("2026-W40") == "Sep 28–Oct 4"
    assert week_range("2026-W53") == "Dec 28–Jan 3"  # across a year boundary
    assert meeting_label("2026-W53") == "Monday, Jan 4"
    assert meeting_label("not a week") == "" and week_range(None) == ""


def test_the_app_opens_on_the_newest_complete_week():
    default_week = load_script("build_site").default_week
    partial, done = {"partial": True}, {"partial": False}
    assert default_week([partial, done, done]) == 1  # Monday: last week's deck, not the new week
    assert default_week([done, done]) == 0
    assert default_week([partial]) == 0
    assert default_week([]) == 0


def test_a_finished_deck_downloads_under_its_monday_name():
    """Vercel names a download after the file, so the .pptx is published as Maya saves it."""
    site = load_script("build_site")
    done = site.week_data(RECORDINGS / "weeks" / "2026-W42" / "run.json")
    assert done["pptx"] == "weeks/2026-W42/competitor-deck-2026-10-19.pptx"
    in_progress = site.week_data(RECORDINGS / "weeks" / "2026-W43" / "run.json")
    assert in_progress["pptx"] is None  # built on Monday, or now with "Build now"


def test_on_monday_the_default_is_last_weeks_deck(tmp_path):
    """Spec 11, manual check 6: once the Monday run has started W44, the app still opens W43."""
    import shutil

    from brightloop.collect.run import build_week

    root = tmp_path / "recordings"
    shutil.copytree(RECORDINGS, root)
    build_week("2026-W43", root)  # the final deck (finished after the week, in this test's time)
    run = json.loads((root / "weeks" / "2026-W43" / "run.json").read_text(encoding="utf-8"))
    run["finished_at"] = "2026-10-26T04:31:00+00:00"
    (root / "weeks" / "2026-W43" / "run.json").write_text(json.dumps(run), encoding="utf-8")
    build_week("2026-W44", root, final=False)
    run = json.loads((root / "weeks" / "2026-W44" / "run.json").read_text(encoding="utf-8"))
    run["finished_at"] = "2026-10-26T04:32:00+00:00"
    (root / "weeks" / "2026-W44" / "run.json").write_text(json.dumps(run), encoding="utf-8")
    site = load_script("build_site")
    weeks = []
    for f in sorted(root.glob("weeks/*/run.json"), reverse=True):
        folder = f.parent
        weeks.append(
            week_view(
                *(
                    json.loads((folder / n).read_text())
                    for n in ("cases.json", "result.json", "run.json")
                )
            )
        )
    assert [w["week"] for w in weeks][:2] == ["2026-W44", "2026-W43"]
    assert weeks[0]["partial"] and not weeks[1]["partial"]
    assert weeks[site.default_week(weeks)]["week"] == "2026-W43"


# ---- Wording edits and page status (spec 5.2) -----------------------------------------------------


def test_an_edit_that_keeps_its_numbers_is_a_wording_edit():
    def edit(old, new):
        return {"kind": "line_changed", "plan": None, "old": old, "new": new}

    assert is_wording_only(edit("1M tokens per month", "1M tokens/month to try Bolt"))
    assert is_wording_only(
        edit("Start at 10M tokens per month", "10M tokens/month, rollover included")
    )
    assert is_wording_only(
        edit("$20 towards most powerful models", "$20 towards more powerful models")
    )
    assert not is_wording_only(edit("$20", "$25"))
    assert not is_wording_only(edit("$20/month", "$20/year"))
    assert not is_wording_only({"kind": "line_removed", "old": "300K tokens daily limit"})


def test_this_weeks_bolt_slide_has_one_change_and_two_wording_edits():
    bolt = page(view("2026-W43"), "Bolt pricing")
    assert bolt["status"] == "changed" and (bolt["real"], bolt["wording"]) == (1, 2)
    assert bolt["lead"] == "Changed this week · 1 change, 2 wording edits"
    rows = [(r["plan"], r["tag"], r["old"], r["new"]) for r in bolt["rows"]]
    assert rows == [
        ("Free", "Removed", "300K tokens daily limit", None),
        ("Free", "Wording", "1M tokens per month", "1M tokens/month to try Bolt"),
        ("Pro", "Wording", "Start at 10M tokens per month", "10M tokens/month, rollover included"),
    ]
    assert bolt["rows"][0]["text"] == 'Free: "300K tokens daily limit" is no longer on the page'


def test_page_statuses_of_the_recorded_weeks():
    w41 = view("2026-W41")
    assert page(w41, "Replit pricing")["status"] == "wording"
    assert page(w41, "Replit homepage")["status"] == "wording"
    assert (w41["changed"], w41["wording"]) == (0, 2)
    assert page(view("2026-W40"), "Bolt pricing")["status"] == "changed"


# ---- The summary slide (spec 4.3.1) ---------------------------------------------------------------


@pytest.mark.parametrize("week", ["2026-W40", "2026-W41", "2026-W42", "2026-W43"])
def test_every_weekly_deck_starts_with_the_summary(week):
    v = view(week)
    summary = v["slides"][0]
    assert summary["status"] == "summary" and len(v["slides"]) == v["pages_total"] + 1 == 7
    assert v["changed"] == sum(s["status"] == "changed" for s in v["slides"][1:])
    assert summary["kicker"] == f"Competitor pages · {week_range(week)}"


def test_the_summary_says_what_changed_and_what_didnt():
    w43 = view("2026-W43")["slides"][0]
    assert w43["title"] == "1 of 6 pages changed"
    first = w43["rows"][0]
    assert (first["type"], first["name"], first["url"]) == (
        "changed",
        "Bolt pricing",
        "https://bolt.new/pricing",
    )
    assert first["more"] == "Plus 2 wording edits"
    assert view("2026-W42")["slides"][0]["title"] == "No changes this week"
    assert view("2026-W41")["slides"][0]["title"] == "No price or limit changed"


# ---- Exceptions and billing (spec 5.3, 6.2) -------------------------------------------------------


def test_archived_copies_and_yearly_prices_are_flagged():
    w42 = view("2026-W42")
    tags = [e["tag"] for e in w42["exceptions"]]
    assert tags[0] == "Archived copies"
    assert w42["exceptions"][0]["text"].startswith(
        "All 6 screenshots come from the Wayback Machine"
    )
    assert not view("2026-W43")["exceptions"]  # all live, monthly


def test_a_yearly_archived_bolt_snapshot_says_so_in_its_footer():
    path = RECORDINGS / "days" / "2026-10-22" / "bolt-pricing.json"
    snap = Snapshot(**json.loads(path.read_text(encoding="utf-8")))
    ref = snapshot_ref(snap, path)
    assert ref["billing"] == "yearly"
    run = {
        "week": "2026-W43",
        "finished_at": "2026-10-26T04:31:00+00:00",
        "pages": [
            {
                "competitor": "Bolt",
                "page": "pricing",
                "url": "https://bolt.new/pricing",
                "status": "unchanged",
                "current": ref,
            }
        ],
    }
    case = {"week": "2026-W43", "timezone": "Asia/Jerusalem", "signals": []}
    summary, bolt = build_slides(case, run_pipeline(case), run)
    assert bolt.footer.endswith(" · archived copy · yearly prices shown")
    assert any("shows yearly prices" in r.get("text", "") for r in summary.rows)


# ---- Links (spec 4.3, 4.4) ----------------------------------------------------------------------


def test_deck_files_link_every_page_to_its_live_page(tmp_path):
    import shutil

    from pptx import Presentation

    folder = tmp_path / "2026-W40"
    shutil.copytree(RECORDINGS / "weeks" / "2026-W40", folder)
    write_deck(folder, assets_root=RECORDINGS)
    html = (folder / "deck.html").read_text(encoding="utf-8")
    assert '<a href="https://bolt.new/pricing"' in html
    assert "<title>Competitor watch · Deck for Monday, Oct 5</title>" in html
    deck = Presentation(str(folder / "deck.pptx"))
    links = []
    for n, slide in enumerate(deck.slides):
        runs = [
            r.hyperlink.address
            for sh in slide.shapes
            if sh.has_text_frame
            for p in sh.text_frame.paragraphs
            for r in p.runs
            if r.hyperlink.address
        ]
        pictures = [sh for sh in slide.shapes if sh.shape_type == 13]
        if n == 0:
            assert "https://bolt.new/pricing" in runs  # the changed page's name
        else:
            assert runs, f"slide {n}: no link on the footer URL"
            assert pictures and all(p.click_action.hyperlink.address for p in pictures)
        links += runs + [p.click_action.hyperlink.address for p in pictures]
    assert links and not any("archive.org" in link for link in links)
    assert "archive.org" not in html


def test_slide_html_escapes_page_text():
    case = {"week": "2026-W43", "timezone": "UTC", "signals": []}
    result = {
        "items": [
            {
                "id": "x",
                "source_ids": [],
                "action": "<script>",
                "change": ['"<b>" → "<i>"'],
                "change_detail": [
                    {"kind": "line_changed", "plan": None, "old": "<b>", "new": "<i>"}
                ],
                "competitor": "Bolt",
                "page": "pricing",
                "links": ["https://bolt.new/pricing"],
            }
        ]
    }
    html = render_html(build_slides(case, result), "t")
    assert "<b>" not in html.split("<body>")[1] and "&lt;b&gt;" in html


# ---- result.json: change_detail (spec 6.1) --------------------------------------------------------


def test_every_item_carries_its_changes_as_data():
    case = json.loads((FIXTURES / "brief_sample.json").read_text(encoding="utf-8"))
    result = run_pipeline(case)
    assert check_contract(result, [s["id"] for s in case["signals"]]) == []
    for week in ("2026-W40", "2026-W41", "2026-W43"):
        result = recorded(week)[1]
        for item in result["items"] + run_pipeline(case)["items"]:
            assert len(item["change_detail"]) == len(item["change"])
            assert all(set(d) == {"kind", "plan", "old", "new"} for d in item["change_detail"])
    bolt = [d for i in recorded("2026-W43")[1]["items"] for d in i["change_detail"]]
    assert {
        "kind": "line_removed",
        "plan": "Free",
        "old": "300K tokens daily limit",
        "new": None,
    } in bolt


# ---- Skip groups (spec 5.4) ---------------------------------------------------------------------


def all_skip_reasons() -> list[str]:
    reasons = []
    files = [FIXTURES / "brief_sample.json", FIXTURES / "edge_cases.json"]
    files += [Path(f) for f in sorted(glob.glob(str(FIXTURES / "adversarial" / "*.json")))]
    for f in files:
        data = json.loads(f.read_text(encoding="utf-8"))
        cases = data if isinstance(data, list) else [data]
        for case in cases:
            if isinstance(case, dict) and "case" in case and "signals" not in case:
                case = case["case"]
            if isinstance(case, dict) and isinstance(case.get("signals"), list):
                reasons += [s["reason"] for s in run_pipeline(case)["skipped"]]
    return reasons


def test_every_skip_reason_has_a_group():
    reasons = all_skip_reasons()
    assert len(reasons) > 50
    assert [r for r in reasons if skip_group(r) == "Other"] == []


def test_the_briefs_sample_groups_as_the_spec_says():
    case = json.loads((FIXTURES / "brief_sample.json").read_text(encoding="utf-8"))
    groups = {s["id"]: skip_group(s["reason"]) for s in run_pipeline(case)["skipped"]}
    assert groups == {
        "sg-303": "Unsafe or unusable",
        "pr-402": "Not from the competitor: verify on the official page",
        "sg-302": "Not from the competitor: verify on the official page",
        "sg-304": "Not from the competitor: verify on the official page",
        "sg-301": "Noa's area: features",
        "ad-117": "Daniel's area: ads",
        "ad-118": "Daniel's area: ads",
    }


# ---- Run log (spec 6.4) -------------------------------------------------------------------------


def test_a_rebuild_adds_exactly_one_run_log_line(tmp_path):
    import shutil

    root = tmp_path / "recordings"
    shutil.copytree(RECORDINGS, root)
    before = (root / "runlog.jsonl").read_text(encoding="utf-8").splitlines()
    entry = load_script("rebuild_history").rebuild(root, "a test")
    after = (root / "runlog.jsonl").read_text(encoding="utf-8").splitlines()
    assert after[: len(before)] == before and len(after) == len(before) + 1
    assert json.loads(after[-1]) == entry
    assert entry["run"] == "rebuild" and entry["reason"] == "a test"
    assert entry["weeks"]["2026-W43"] == {"signals": 3, "items": 3}


def test_daily_runs_list_their_errors():
    lines = (RECORDINGS / "runlog.jsonl").read_text(encoding="utf-8").splitlines()
    daily = [json.loads(line) for line in lines if '"run": "daily"' in line]
    assert daily and "errors" in daily[-1]


def test_the_run_log_table_shows_the_last_14_runs():
    site = load_script("build_site")
    rows = site.run_rows(site.read_runlog())
    assert 0 < len(rows) <= 14
    assert all(r["run"] in ("Daily capture", "Weekly deck", "Rebuild", "Delivery") for r in rows)
    assert all(r["pages"] is None or len(r["pages"]) == 6 for r in rows)


# ---- The deck is built on Monday; the week in progress on request (decision D19) ---------------


def test_only_the_monday_deck_writes_deck_files(tmp_path):
    import shutil

    from brightloop.collect.run import build_week

    root = tmp_path / "recordings"
    shutil.copytree(RECORDINGS, root)
    folder = root / "weeks" / "2026-W43"
    build_week("2026-W43", root, final=False)  # the daily refresh of the week so far
    assert (folder / "run.json").exists()
    assert not (folder / "deck.pptx").exists() and not (folder / "deck.html").exists()
    build_week("2026-W43", root)  # Monday's final deck
    assert (folder / "deck.pptx").exists() and (folder / "deck.html").exists()


def test_build_deck_now_makes_the_files_on_request():
    spec = importlib.util.spec_from_file_location("api_deck", ROOT / "api" / "deck.py")
    api = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(api)
    body, kind, name = api.deck_file("2026-W43", RECORDINGS, "pptx")
    assert body[:2] == b"PK" and kind.endswith("presentationml.presentation")
    assert name.startswith("competitor-deck-2026-10-26") and name.endswith(".pptx")
    body, kind, name = api.deck_file("2026-W43", RECORDINGS, "html")
    html = body.decode("utf-8")
    assert kind.startswith("text/html") and '<a href="https://bolt.new/pricing"' in html
    assert 'src="/days/' in html and "archive.org" not in html


def test_the_backup_schedule_recognises_todays_scheduled_run():
    """The workflow's backup schedules stop when today's scheduled run is in the run log."""
    import re

    workflow = (ROOT / ".github" / "workflows" / "daily.yml").read_text(encoding="utf-8")
    assert 'grep -q "\\"date\\": \\"$today\\".*\\"trigger\\": \\"scheduled\\""' in workflow
    lines = (RECORDINGS / "runlog.jsonl").read_text(encoding="utf-8").splitlines()
    scheduled = [line for line in lines if '"run": "daily"' in line and '"scheduled"' in line]
    assert scheduled
    day = json.loads(scheduled[-1])["date"]
    assert re.search(f'"date": "{day}".*"trigger": "scheduled"', scheduled[-1])
