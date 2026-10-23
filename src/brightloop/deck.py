"""Maya's Monday deck: one slide per competitor page, built from a week's run + result.

The deck is a wrapper for leadership: short and clean. Each slide is the page's screenshot,
one status line and, when the analysis kept something, one dated line per item (the item's
`action`, as the pipeline wrote it). No confidence, sources or caveats on the slides: those
stay in result.json. Wrong content on a slide is a problem in an earlier step, not here.

Rendered as HTML (the web app) and .pptx (opens in Google Slides via File → Import), from the
same slide data. Design reference: review/deck-mock/slide-v3.html.
"""

from __future__ import annotations

import html
import json
import os
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

PAGES = (("pricing", "pricing"), ("home", "homepage"))
DEFAULT_COMPETITORS = ("Lovable", "Replit", "Bolt")
STATUS_LINE = {
    "changed": "Changed this week",
    "unchanged": "No change this week",
    "not-captured": "Not captured this week",
    "first-snapshot": "First capture: nothing to compare yet",
    "failed-capture": "Capture failed: check the page by hand",
}
BRAND = "Competitor watch"


@dataclass
class Slide:
    title: str  # "Bolt pricing"
    status: str  # changed | unchanged | not-captured | first-snapshot | failed-capture
    lead: str  # one status line
    rows: list[tuple[str, str]] = field(default_factory=list)  # (date, item action)
    url: str | None = None
    shot: str | None = None  # screenshot path
    footer: str = ""


def build_slides(
    case: dict[str, Any], result: dict[str, Any], run: dict[str, Any] | None = None
) -> list[Slide]:
    """With the collector's run.json: one slide per page, every page shown (also "no change").
    Without it (an uploaded cases.json): one slide per item."""
    tz = ZoneInfo(case.get("timezone") or "UTC")
    week = week_label(case.get("week"))
    signal_ts = {s.get("id"): s.get("ts") for s in case.get("signals", [])}
    item_of = {sid: item for item in result.get("items", []) for sid in item["source_ids"]}
    entries = sorted(
        (run or {}).get("pages", []),
        key=lambda e: (*order(e["competitor"]), [p for p, _ in PAGES].index(e["page"])),
    )
    slides, shown = [], set()
    for entry in entries:
        items = unique_items(item_of.get(sid) for sid in entry.get("signal_ids", []))
        shown.update(item["id"] for item in items)
        slides.append(page_slide(entry, items, signal_ts, week, tz))
    for item in result.get("items", []):
        if item["id"] not in shown:
            slides.append(item_slide(item, signal_ts, week, tz))
    return slides


def page_slide(
    entry: dict[str, Any],
    items: list[dict[str, Any]],
    signal_ts: dict[str, str],
    week: str,
    tz: ZoneInfo,
) -> Slide:
    status = entry["status"]
    if status == "changed" and not items:
        status = "unchanged"  # the page differed, but the analysis kept nothing for Maya
    current = entry.get("current") or {}
    return Slide(
        title=f"{entry['competitor']} {dict(PAGES)[entry['page']]}",
        status=status,
        lead=STATUS_LINE.get(status, status),
        rows=[(item_date(item, signal_ts, tz), item["action"]) for item in items],
        url=entry["url"],
        shot=current.get("screenshot") if status != "not-captured" else None,
        footer=footer(week, entry["url"], current, tz),
    )


def item_slide(item: dict[str, Any], signal_ts: dict[str, str], week: str, tz: ZoneInfo) -> Slide:
    page = dict(PAGES).get(item.get("page", ""), "")
    url = (item.get("links") or [None])[0] or (item.get("cta") or {}).get("url")
    return Slide(
        title=f"{item.get('competitor', '')} {page}".strip(),
        status="changed",
        lead=STATUS_LINE["changed"],
        rows=[(item_date(item, signal_ts, tz), item["action"])],
        url=url,
        footer=footer(week, url, {}, tz),
    )


def unique_items(items) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for item in items:
        if item and all(item["id"] != o["id"] for o in out):
            out.append(item)
    return out


def item_date(item: dict[str, Any], signal_ts: dict[str, str], tz: ZoneInfo) -> str:
    """The day the item's first signal was seen, e.g. "Sep 30"."""
    stamps = sorted(ts for sid in item["source_ids"] if (ts := signal_ts.get(sid)))
    return short_date(stamps[0], tz) if stamps else ""


def footer(week: str, url: str | None, current: dict[str, Any], tz: ZoneInfo) -> str:
    parts = [week, (url or "").removeprefix("https://").removeprefix("www.").rstrip("/")]
    if current.get("captured_at"):
        when = short_date(current["captured_at"], tz)
        archived = current.get("method") == "archived"
        parts.append(f"screenshot {when}" + (" (archived copy)" if archived else ""))
    return " · ".join(p for p in parts if p)


# ---- Rendering ----------------------------------------------------------------------------


def render_html(slides: list[Slide], title: str) -> str:
    def esc(value: str | None) -> str:
        return html.escape(value or "")

    parts = []
    for s in slides:
        rows = "".join(
            f"<li><time>{esc(day)}</time><span>{esc(text)}</span></li>" for day, text in s.rows
        )
        shot = (
            f'<div class="shot"><img src="{esc(s.shot)}" alt="{esc(s.title)}"></div>'
            if s.shot
            else ""
        )
        parts.append(
            f'<div class="frame"><section class="slide {s.status}"><div class="text">'
            f'<h2>{esc(s.title)}</h2><p class="lead">{esc(s.lead)}</p>'
            f"{'<ul>' + rows + '</ul>' if rows else ''}</div>{shot}"
            f"<footer><span>{esc(s.footer)}</span><span>{BRAND}</span></footer></section></div>"
        )
    return HTML_PAGE.replace("{{title}}", html.escape(title)).replace(
        "{{slides}}", "\n".join(parts)
    )


# Sizes are in slide pixels (a 1920×1080 slide): calc(N * var(--u)) scales them to the window.
HTML_PAGE = """<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>{{title}}</title>
<link rel="preconnect" href="https://fonts.googleapis.com">
<link href="https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600&display=swap" rel="stylesheet">
<style>
*{box-sizing:border-box}
body{margin:0;background:#e4e4e7;padding:24px 16px;font-family:Inter,system-ui,-apple-system,"Segoe UI",sans-serif;
  color:#18181b;-webkit-font-smoothing:antialiased}
.frame{max-width:1280px;margin:0 auto 24px;container-type:inline-size}
.slide{--u:calc(100cqw / 1920);aspect-ratio:16/9;background:#fff;overflow:hidden;
  border-radius:calc(8 * var(--u));box-shadow:0 1px 3px rgba(0,0,0,.08);
  padding:calc(96 * var(--u)) calc(104 * var(--u)) calc(64 * var(--u));
  display:grid;grid-template-columns:calc(560 * var(--u)) 1fr;grid-template-rows:1fr auto;column-gap:calc(88 * var(--u))}
.text{display:flex;flex-direction:column;min-width:0}
h2{margin:0;font-size:calc(72 * var(--u));font-weight:600;letter-spacing:-.025em;line-height:1.05}
.lead{margin:calc(28 * var(--u)) 0 0;font-size:calc(32 * var(--u));line-height:1.4;color:#52525b}
ul{list-style:none;margin:calc(64 * var(--u)) 0 0;padding:0;display:flex;flex-direction:column;min-height:0;overflow:hidden}
li{display:grid;grid-template-columns:calc(110 * var(--u)) 1fr;padding:calc(22 * var(--u)) 0;
  border-top:1px solid #e4e4e7;font-size:calc(26 * var(--u));line-height:1.4}
li:last-child{border-bottom:1px solid #e4e4e7}
li time{color:#a1a1aa;font-weight:500}
.shot{border-radius:calc(14 * var(--u));overflow:hidden;border:1px solid #e4e4e7;align-self:start;
  box-shadow:0 calc(30 * var(--u)) calc(60 * var(--u)) calc(-30 * var(--u)) rgba(0,0,0,.28);
  max-height:calc(820 * var(--u))}
.shot img{width:100%;display:block}
footer{grid-column:1/-1;display:flex;justify-content:space-between;gap:calc(24 * var(--u));
  font-size:calc(18 * var(--u));color:#a1a1aa;padding-top:calc(36 * var(--u))}
@media print{body{background:#fff;padding:0}.frame{max-width:none;margin:0;break-after:page}.slide{border-radius:0;box-shadow:none}}
</style></head><body>
{{slides}}
</body></html>
"""


def render_pptx(slides: list[Slide], folder: Path, path: Path) -> None:
    """Same layout as the HTML slide: text on the left, the screenshot on the right."""
    from PIL import Image
    from pptx import Presentation  # optional dependency: pip install .[deck]
    from pptx.dml.color import RGBColor
    from pptx.util import Emu, Pt

    px = 12192000 // 1920  # EMU per slide pixel (1920×1080 layout, as in the HTML)
    ink, soft, faint = (
        RGBColor(0x18, 0x18, 0x1B),
        RGBColor(0x52, 0x52, 0x5B),
        RGBColor(0xA1, 0xA1, 0xAA),
    )
    deck = Presentation()
    deck.slide_width, deck.slide_height = Emu(1920 * px), Emu(1080 * px)
    blank = deck.slide_layouts[6]

    def text(slide, x, y, w, h, runs):
        frame = slide.shapes.add_textbox(
            Emu(x * px), Emu(y * px), Emu(w * px), Emu(h * px)
        ).text_frame
        frame.word_wrap = True
        for n, (value, size, color, bold) in enumerate(runs):
            para = frame.paragraphs[0] if n == 0 else frame.add_paragraph()
            run = para.add_run()
            run.text = value
            run.font.size, run.font.bold = Pt(size), bold
            run.font.color.rgb = color
            para.space_after = Pt(size * 0.6)

    for s in slides:
        slide = deck.slides.add_slide(blank)
        rows = [(f"{day}   {action}" if day else action, 14, ink, False) for day, action in s.rows]
        text(slide, 104, 96, 560, 820, [(s.title, 36, ink, True), (s.lead, 18, soft, False), *rows])
        text(slide, 104, 980, 1712, 40, [(f"{s.footer}    {BRAND}", 10, faint, False)])
        if s.shot and (folder / s.shot).exists():
            with Image.open(folder / s.shot) as img:
                ratio = img.width / img.height
            w = 1920 - 104 - (104 + 560 + 88)
            h = min(int(w / ratio), 820)
            left, top = 104 + 560 + 88, 96
            pic = slide.shapes.add_picture(
                str(folder / s.shot), Emu(left * px), Emu(top * px), Emu(w * px)
            )
            if int(w / ratio) > h:  # crop the bottom of tall screenshots to the slide area
                pic.crop_bottom = 1 - h / (w / ratio)
                pic.height = Emu(h * px)
    deck.save(str(path))


def write_deck(folder: Path, pptx: bool = True, assets_root: Path | None = None) -> list[Path]:
    """Build deck.html (and deck.pptx) next to a run's cases.json and result.json.

    Screenshot paths in run.json are relative to `assets_root` (recordings/); the deck
    rewrites them relative to its own folder so the HTML works when served as files."""
    case = json.loads((folder / "cases.json").read_text(encoding="utf-8"))
    result = json.loads((folder / "result.json").read_text(encoding="utf-8"))
    run_path = folder / "run.json"
    run = json.loads(run_path.read_text(encoding="utf-8")) if run_path.exists() else None
    slides = build_slides(case, result, run)
    if assets_root is not None:
        prefix = Path(os.path.relpath(assets_root, folder)).as_posix()
        for slide in slides:
            if slide.shot:
                slide.shot = f"{prefix}/{slide.shot}"
    title = f"{BRAND} · {week_label(case.get('week')) or folder.name}"
    out = [folder / "deck.html"]
    out[0].write_text(render_html(slides, title), encoding="utf-8", newline="\n")
    if pptx:
        try:
            render_pptx(slides, folder, folder / "deck.pptx")
            out.append(folder / "deck.pptx")
        except ImportError:
            pass  # .pptx needs python-pptx; the HTML deck is always written
    return out


# ---- Helpers ---------------------------------------------------------------------------------


def week_label(week: str | None) -> str:
    """'2026-W40' -> 'Week of Sep 28'."""
    try:
        year, number = week.split("-W")
        monday = date.fromisocalendar(int(year), int(number), 1)
    except (AttributeError, ValueError):
        return ""
    return f"Week of {monday:%b} {monday.day}"


def short_date(ts: str, tz: ZoneInfo) -> str:
    try:
        moment = datetime.fromisoformat(ts).astimezone(tz)
    except ValueError:
        return ts
    return f"{moment:%b} {moment.day}"


def order(competitor: str) -> tuple[int, str]:
    known = list(DEFAULT_COMPETITORS)
    return (known.index(competitor) if competitor in known else len(known), competitor)
